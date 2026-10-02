# =====================================================================
#  EMA 5-8-13 BULLISH CROSSOVER SCANNER  (Daily | NSE | Upstox | Colab)
# =====================================================================
#  SETUP (ek baar):
#   1) Colab Secrets (left side key icon) me ye 3 naam bana, "Notebook access" ON:
#        UPSTOX_TOKEN  ->  Upstox access token (roz naya, expire hota hai)
#        TG_TOKEN      ->  Telegram bot token
#        TG_CHAT_ID    ->  Telegram chat id
#      (tere naam alag hain to neeche CONFIG me badal de)
#   2) Google Drive me folder bana:  MyDrive/EMA_Scanner/  aur usme stocklist.xlsx rakh
#   3) Roz EOD (4:00 PM IST ke baad) -> Runtime > Run all
#
#  OUTPUT (sab Drive ke EMA_Scanner folder me):
#   EMA_Scanner_Master.xlsx : Signals | Volume Spike | Sector Daily | Sector Weekly
#   Sector_Momentum.xlsx    : Latest Ranking | Score Daily | Score Weekly
# =====================================================================
import os, io, re, gzip, json, time, html, threading
from datetime import datetime, timedelta, time as dtime
from zoneinfo import ZoneInfo
from urllib.parse import quote
from concurrent.futures import ThreadPoolExecutor

import requests
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================ CONFIG ================================
BASE_DIR      = os.environ.get("EMA_BASE_DIR", "/content/drive/MyDrive/EMA_Scanner")
STOCKLIST     = f"{BASE_DIR}/stocklist.xlsx"
MASTER_XLSX   = f"{BASE_DIR}/EMA_Scanner_Master.xlsx"
MOMENTUM_XLSX = f"{BASE_DIR}/Sector_Momentum.xlsx"
UNIVERSE_CSV  = f"{BASE_DIR}/liquid_universe.csv"   # liquid stocks ki list (auto banti hai)
SECTOR_LOG    = f"{BASE_DIR}/sector_log.csv"        # sector history (auto banti hai)

SECRET_UPSTOX, SECRET_TG_TOKEN, SECRET_TG_CHAT = "UPSTOX_TOKEN", "TG_TOKEN", "TG_CHAT_ID"

# --- Signal rules ---
EMA_FAST, EMA_MID, EMA_SLOW = 5, 8, 13
CROSS_PAIRS      = [(5, 8), (5, 13), (8, 13)]  # inme se koi bhi bullish cross (fast > slow)
CROSS_LOOKBACK   = 3      # pichle kitne trading din me cross hua ho (1 = sirf aaj ka fresh cross)
MIN_EMAS_BELOW   = 2      # close kam se kam itni EMAs ke upar ho (3 me se)
TOP_N            = 50     # roz Signals sheet me top kitne
RANK_BY          = "vol_ratio"   # "vol_ratio" (aaj ka vol / 20d avg) ya "volume" (raw volume)

# --- Liquidity filter ---
MIN_PRICE             = 20
MIN_TRADED_VALUE_CR   = 1.0     # 20-day avg (price x volume) kam se kam Rs 1 Cr
LIQ_DAYS              = 20
UNIVERSE_REFRESH_DAYS = 7       # itne din baad poori list dobara scan hogi
FORCE_UNIVERSE_REFRESH = False  # True = aaj hi poori list dobara scan karo

# --- Volume spike ---
VOL_SPIKE_X   = 3.0     # aaj ka volume >= 3x (20d avg) AND green close
VOL_SPIKE_TOP = 50

# --- Sector ---
MIN_SECTOR_STOCKS = 5   # isse kam stocks wale sector ka score nahi banega
KEEP_DAILY_COLS   = 90
KEEP_WEEKLY_COLS  = 26

# --- API / misc ---
HISTORY_DAYS    = 130   # calendar days ka history (EMA warm-up + 20d avg ke liye kaafi)
WORKERS         = 5
MIN_GAP_SEC     = 0.12  # 2 requests ke beech minimum gap (rate limit safe)
SEND_EXCEL_TO_TG = True
IST = ZoneInfo("Asia/Kolkata")

# ======================== SECTOR MAPPING ============================
# Industry -> broad sector (upar se neeche, pehla match jeetega)
SECTOR_RULES = [
    ("FMCG & Food",                    r"edible"),
    ("Realty",                         r"realty|real estate|residential|reit"),
    ("Financial Services",             r"financ|nbfc|non banking|insurance|asset manag|investment|holding|broking|capital market"),
    ("Banks",                          r"bank"),
    ("IT",                             r"^it\b|software|bpo|kpo|business process|outsourcing"),
    ("Healthcare & Pharma",            r"pharma|biotech|health|hospital|medical"),
    ("Auto",                           r"auto|tyre|wheeler|vehicle"),
    ("Chemicals & Agro",               r"chemical|carbon black|industrial gas|^paints|fertili|pesticide"),
    ("Oil & Gas",                      r"oil|refiner|petro|exploration|offshore|drilling|gas"),
    ("Power & Utilities",              r"electric utilit|power|utilit|waste"),
    ("Telecom",                        r"telecom"),
    ("Metals & Mining",                r"iron|steel|alumin|copper|zinc|mining|coal|non-ferrous|precious metal|castings"),
    ("Cement & Building Materials",    r"cement|construction materials"),
    ("Infrastructure & Construction",  r"construction|civil|roads|engineering"),
    ("Textiles & Apparel",             r"textile|apparel|footwear|jute|fibre"),
    ("Retail & E-commerce",            r"retail|department store|e-commerce|gems|jewel|gift|catalogue"),
    ("Consumer Durables & Electronics", r"consumer electronic|appliance|houseware|furniture|electronic component|computer hardware|storage media|photographic"),
    ("FMCG & Food",                    r"personal product|household|food|sugar|\btea\b|dairy|beverage|brewer|distiller|cigarette|tobacco|agricultural|packaged"),
    ("Media & Entertainment",          r"movie|film|broadcast|publish|publication|advertis|entertainment"),
    ("Hospitality & Travel",           r"hotel|restaurant|leisure|travel|tour"),
    ("Logistics & Transport",          r"transport|shipping|logistic|marine|airline|surface"),
    ("Paper, Packaging & Plastics",    r"paper|packag|container|plastic|forest|station|printing"),
    ("Capital Goods",                  r"machinery|heavy electrical|industrial|electrical|elect\.|equipment|defen[cs]e|aerospace"),
    ("Education & Business Services",  r"education|e-learning|training|commercial service|consulting|sp\.consumer|service"),
    ("Trading & Distribution",         r"trading|distribut"),
    ("Diversified & Others",           r"diversified|others|unknown"),
]

def map_sector(industry):
    s = html.unescape(str(industry or "")).strip().lower()
    s = re.sub(r"\s+", " ", s)
    if s in ("", "nan", "none"):
        return "Unclassified"
    if s == "sme":
        return "SME"
    for sector, pat in SECTOR_RULES:
        if re.search(pat, s):
            return sector
    return "Diversified & Others"

# ========================== HELPERS =================================
def get_secret(name):
    try:
        from google.colab import userdata
        return userdata.get(name)
    except Exception:
        return os.environ.get(name)

def mount_drive():
    try:
        from google.colab import drive
        drive.mount("/content/drive")
    except Exception:
        pass
    os.makedirs(BASE_DIR, exist_ok=True)

def norm(sym):
    return re.sub(r"[^A-Z0-9]", "", str(sym).upper())

def py(v):
    if v is None:
        return None
    if isinstance(v, (np.floating, float)):
        return None if np.isnan(v) else float(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.bool_):
        return bool(v)
    if isinstance(v, pd.Timestamp):
        return v.date()
    return v

# ===================== TELEGRAM =====================================
TG = {"token": None, "chat": None}

def tg_send(text):
    if not (TG["token"] and TG["chat"]):
        print("[Telegram skipped - token/chat id nahi mila]\n" + text)
        return
    chunks = [text[i:i + 3900] for i in range(0, len(text), 3900)]
    for ch in chunks:
        try:
            requests.post(f"https://api.telegram.org/bot{TG['token']}/sendMessage",
                          json={"chat_id": TG["chat"], "text": ch, "parse_mode": "HTML",
                                "disable_web_page_preview": True}, timeout=30)
        except Exception as e:
            print("Telegram error:", e)

def tg_doc(path, caption=""):
    if not (TG["token"] and TG["chat"]) or not os.path.exists(path):
        return
    try:
        with open(path, "rb") as f:
            requests.post(f"https://api.telegram.org/bot{TG['token']}/sendDocument",
                          data={"chat_id": TG["chat"], "caption": caption},
                          files={"document": f}, timeout=120)
    except Exception as e:
        print("Telegram doc error:", e)

# ===================== UPSTOX API ===================================
class Limiter:
    def __init__(self, gap):
        self.gap, self.last, self.lock = gap, 0.0, threading.Lock()
    def wait(self):
        with self.lock:
            now = time.time()
            delta = self.last + self.gap - now
            if delta > 0:
                time.sleep(delta)
            self.last = time.time()

LIMITER = Limiter(MIN_GAP_SEC)
HEADERS = {}
AUTH_FAILED = {"flag": False}

def api_get(url):
    for attempt in range(6):
        LIMITER.wait()
        try:
            r = requests.get(url, headers=HEADERS, timeout=30)
        except requests.RequestException:
            time.sleep(2 * (attempt + 1)); continue
        if r.status_code == 200:
            return r.json()
        if r.status_code == 401:
            AUTH_FAILED["flag"] = True
            return None
        if r.status_code == 429:            # rate limit -> ruk ke retry
            time.sleep(15 * (attempt + 1)); continue
        if r.status_code >= 500:
            time.sleep(2 * (attempt + 1)); continue
        return None
    return None

def candles_to_df(j):
    try:
        cs = j["data"]["candles"]
    except Exception:
        return None
    if not cs:
        return None
    df = pd.DataFrame([c[:6] for c in cs], columns=["ts", "open", "high", "low", "close", "volume"])
    df["date"] = pd.to_datetime(df["ts"].astype(str).str[:10]).dt.date
    return df.drop(columns="ts").sort_values("date").reset_index(drop=True)

def fetch_daily(key, to_date, from_date):
    url = f"https://api.upstox.com/v3/historical-candle/{quote(key, safe='')}/days/1/{to_date}/{from_date}"
    return candles_to_df(api_get(url))

def fetch_today_intraday(key):
    url = f"https://api.upstox.com/v3/historical-candle/intraday/{quote(key, safe='')}/days/1"
    return candles_to_df(api_get(url))

def check_token():
    j = api_get("https://api.upstox.com/v2/user/profile")
    return j is not None

def load_instruments():
    r = requests.get("https://assets.upstox.com/market-quote/instruments/exchange/NSE.json.gz", timeout=90)
    r.raise_for_status()
    ins = pd.DataFrame(json.loads(gzip.decompress(r.content)))
    ins = ins[(ins["segment"] == "NSE_EQ") & (ins["instrument_type"] == "EQ")]
    ins = ins[["trading_symbol", "instrument_key"]].copy()
    ins["k"] = ins["trading_symbol"].map(norm)
    return ins.drop_duplicates("k")

# ===================== STOCKLIST ====================================
def load_stocklist():
    raw = pd.read_excel(STOCKLIST, header=None, dtype=str)
    hdr = None
    for i in range(min(10, len(raw))):
        vals = [str(x).strip().upper() for x in raw.iloc[i].tolist()]
        if "NSE" in vals:
            hdr = i; break
    if hdr is None:
        raise ValueError("stocklist.xlsx me 'NSE' column header nahi mila")
    df = raw.iloc[hdr + 1:].copy()
    df.columns = [str(c).strip().upper() for c in raw.iloc[hdr]]
    name_col = next((c for c in df.columns if "NAME" in c), df.columns[0])
    ind_col = next((c for c in df.columns if "INDUSTRY" in c), None)
    out = pd.DataFrame({
        "symbol": df["NSE"].astype(str).str.strip().str.upper(),
        "name": df[name_col].astype(str).str.strip(),
        "industry": df[ind_col].fillna("").astype(str).map(lambda x: html.unescape(x).strip()) if ind_col else "",
    })
    out = out[~out["symbol"].isin(["", "NAN", "NONE"])].drop_duplicates("symbol")
    out["sector"] = out["industry"].map(map_sector)
    out["k"] = out["symbol"].map(norm)
    return out.reset_index(drop=True)

# ===================== ANALYSIS =====================================
def analyse(df, liq_days=LIQ_DAYS):
    """Ek stock ka sab calculation. df: date/open/high/low/close/volume (ascending)"""
    if df is None or len(df) < 30:
        return None
    c = df["close"].astype(float).reset_index(drop=True)
    v = df["volume"].astype(float).reset_index(drop=True)
    ema = {n: c.ewm(span=n, adjust=False).mean() for n in {EMA_FAST, EMA_MID, EMA_SLOW}}
    last = len(c) - 1
    close, prev = c.iloc[last], c.iloc[last - 1]
    avg_vol = v.iloc[-21:-1].mean()
    tv_cr = (c.iloc[-liq_days:] * v.iloc[-liq_days:]).mean() / 1e7

    cross_k, cross_txt = None, ""
    for k in range(CROSS_LOOKBACK):
        i = last - k
        hit = [f"{a}>{b}" for a, b in CROSS_PAIRS
               if ema[a].iloc[i] > ema[b].iloc[i] and ema[a].iloc[i - 1] <= ema[b].iloc[i - 1]]
        if hit:
            cross_k, cross_txt = k, ", ".join(hit); break

    e_f, e_m, e_s = ema[EMA_FAST].iloc[last], ema[EMA_MID].iloc[last], ema[EMA_SLOW].iloc[last]
    above = int(close > e_f) + int(close > e_m) + int(close > e_s)
    signal = (cross_k is not None) and (e_f > e_s) and (above >= MIN_EMAS_BELOW)
    return {
        "close": close, "chg_pct": (close / prev - 1) * 100 if prev else np.nan,
        "ema_f": e_f, "ema_m": e_m, "ema_s": e_s,
        "cross": cross_txt, "cross_k": cross_k, "above_n": above,
        "aligned": bool(e_f > e_m > e_s),
        "volume": v.iloc[last], "avg_vol": avg_vol,
        "vol_ratio": (v.iloc[last] / avg_vol) if avg_vol and avg_vol > 0 else np.nan,
        "tv_cr": tv_cr, "signal": bool(signal),
    }

def cross_tag(k):
    if k is None or (isinstance(k, float) and np.isnan(k)):
        return ""
    return "FRESH" if int(k) == 0 else f"{int(k)}d ago"

# ===================== EXCEL: STYLING ===============================
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
GREEN_F, RED_F = Font(color="006100"), Font(color="9C0006")
TOP_BORDER = Border(top=Side(style="medium", color="1F3864"))

def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def get_wb(path):
    return load_workbook(path) if os.path.exists(path) else Workbook()

def save_wb(wb, path, order):
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + \
                 [wb[n] for n in wb.sheetnames if n not in order]
    wb.active = 0
    wb.save(path)

def append_sheet(wb, name, df, run_date, fmts=None, widths=None):
    """Same sheet me roz data niche append. Same date dobara chale to purana usi date ka data replace."""
    cols = list(df.columns)
    if name not in wb.sheetnames or wb[name].max_row < 1 or wb[name]["A1"].value is None:
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name)
        ws.append(cols)
        style_header(ws)
        ws.freeze_panes = "A2"
        for i, cn in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(cn, max(11, len(cn) + 3))
    else:
        ws = wb[name]
    for r in range(ws.max_row, 1, -1):
        v = ws.cell(r, 1).value
        v = v.date() if isinstance(v, datetime) else v
        if v == run_date:
            ws.delete_rows(r)
    start = ws.max_row + 1
    for row in df.itertuples(index=False):
        ws.append([py(x) for x in row])
    for r in range(start, ws.max_row + 1):
        for i, cn in enumerate(cols, 1):
            cell = ws.cell(r, i)
            if cn == "Date":
                cell.number_format = "dd-mmm-yyyy"
            elif fmts and cn in fmts:
                cell.number_format = fmts[cn]
            if r == start:
                cell.border = TOP_BORDER
            if cn == "Signal" and cell.value == "FRESH":
                cell.fill, cell.font = GREEN, GREEN_F
    ws.auto_filter.ref = ws.dimensions
    return ws

def write_heatmap(wb, name, mat, fmt="0", note=""):
    """mat: index=sector, columns=period label (newest -> oldest). Color: green = pichle se zyada, red = kam."""
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    labels = list(mat.columns)
    ws.append(["Sector", "Change vs prev"] + labels)
    style_header(ws)
    ws.row_dimensions[1].height = 32
    m = mat.sort_values(by=labels[0], ascending=False, na_position="last") if labels else mat
    for sector, row in m.iterrows():
        vals = [None if pd.isna(x) else float(x) for x in row.tolist()]
        delta = (vals[0] - vals[1]) if len(vals) > 1 and vals[0] is not None and vals[1] is not None else None
        ws.append([sector, delta] + [("-" if x is None else x) for x in vals])
        r = ws.max_row
        for j in range(len(vals)):
            cell = ws.cell(r, 3 + j)
            cell.number_format, cell.alignment = fmt, Alignment(horizontal="center")
            nxt = vals[j + 1] if j + 1 < len(vals) else None
            if vals[j] is not None and nxt is not None:
                if vals[j] > nxt:
                    cell.fill, cell.font = GREEN, GREEN_F
                elif vals[j] < nxt:
                    cell.fill, cell.font = RED, RED_F
        dc = ws.cell(r, 2)
        dc.number_format, dc.alignment = "+0.0;-0.0;0.0", Alignment(horizontal="center")
        if delta is not None and delta != 0:
            dc.fill, dc.font = (GREEN, GREEN_F) if delta > 0 else (RED, RED_F)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14
    for j in range(len(labels)):
        ws.column_dimensions[get_column_letter(3 + j)].width = 13
    ws.freeze_panes = "C2"
    ws.cell(ws.max_row + 2, 1, "Green = pichle period se zyada | Red = kam | " + note).font = Font(italic=True, color="595959")

# ===================== SECTOR LOG / MATRICES ========================
def update_sector_log(res, run_date):
    """Roz ke sector stats ek log me store (sab heatmap/score isi se bante hain)."""
    rows = []
    for sec, g in res.groupby("sector"):
        total, sig = len(g), int(g["signal"].sum())
        spikes = int(g["spike"].sum())
        sg = g[g["signal"]]
        breadth = sig / total * 100 if total else np.nan
        volconf = (sg["vol_ratio"] >= 1.5).mean() * 100 if sig else 0.0
        score = (0.7 * breadth + 0.3 * volconf) if total >= MIN_SECTOR_STOCKS else np.nan
        rows.append({"Date": run_date, "Sector": sec, "Total": total, "Signals": sig, "Spikes": spikes,
                     "Breadth": round(breadth, 2), "VolConfirm": round(volconf, 2),
                     "Score": round(score, 2) if not np.isnan(score) else np.nan,
                     "AvgVolRatio": round(sg["vol_ratio"].mean(), 2) if sig else np.nan})
    new = pd.DataFrame(rows)
    if os.path.exists(SECTOR_LOG):
        old = pd.read_csv(SECTOR_LOG)
        old["Date"] = pd.to_datetime(old["Date"]).dt.date
        old = old[old["Date"] != run_date]
        new = pd.concat([old, new], ignore_index=True)
    new = new.sort_values(["Date", "Sector"]).reset_index(drop=True)
    new.to_csv(SECTOR_LOG, index=False)
    return new

def daily_matrix(log, col):
    p = log.pivot_table(index="Sector", columns="Date", values=col, aggfunc="first", dropna=False)
    p = p[sorted(p.columns, reverse=True)[:KEEP_DAILY_COLS]]
    p.columns = [d.strftime("%d-%b-%y") for d in p.columns]
    return p

def weekly_matrix(log, col):
    d = log.copy()
    d["wk"] = pd.to_datetime(d["Date"]).dt.strftime("%G-W%V")
    last_day = d.groupby("wk")["Date"].max()
    p = d.groupby(["Sector", "wk"])[col].mean().unstack()
    wks = sorted(p.columns, reverse=True)[:KEEP_WEEKLY_COLS]
    p = p[wks]
    p.columns = [f"{w}\n(till {last_day[w]:%d-%b})" for w in wks]
    return p

def write_latest_ranking(wb, log, run_date):
    d = log[log["Date"] == run_date].copy()
    dates = sorted(log["Date"].unique())
    prev_date = dates[-2] if len(dates) > 1 and dates[-1] == run_date else None
    prev = log[log["Date"] == prev_date].set_index("Sector")["Score"] if prev_date else pd.Series(dtype=float)
    wk = weekly_matrix(log, "Score")
    d = d.dropna(subset=["Score"]).sort_values("Score", ascending=False)
    if "Latest Ranking" in wb.sheetnames:
        del wb["Latest Ranking"]
    ws = wb.create_sheet("Latest Ranking")
    ws.append(["Rank", "Sector", "Score (0-100)", "Change vs prev day", "Weekly avg score", "Change vs prev week",
               "Stocks scanned", "EMA signals", "Breadth %", "Vol-confirm %", "Vol spikes"])
    style_header(ws); ws.row_dimensions[1].height = 32
    for i, r in enumerate(d.itertuples(index=False), 1):
        dd = r.Score - prev.get(r.Sector, np.nan) if r.Sector in prev.index and not pd.isna(prev.get(r.Sector)) else None
        w0 = w1 = None
        if r.Sector in wk.index:
            vals = wk.loc[r.Sector].tolist()
            w0 = None if pd.isna(vals[0]) else vals[0]
            w1 = vals[1] if len(vals) > 1 and not pd.isna(vals[1]) else None
        wd = (w0 - w1) if (w0 is not None and w1 is not None) else None
        ws.append([i, r.Sector, r.Score, dd, w0, wd, r.Total, r.Signals, r.Breadth, r.VolConfirm, r.Spikes])
        row = ws.max_row
        for col, f in ((3, "0.0"), (4, "+0.0;-0.0;0.0"), (5, "0.0"), (6, "+0.0;-0.0;0.0"), (9, "0.0"), (10, "0.0")):
            ws.cell(row, col).number_format = f
        for col in (4, 6):
            v = ws.cell(row, col).value
            if v is not None and v != 0:
                ws.cell(row, col).fill, ws.cell(row, col).font = (GREEN, GREEN_F) if v > 0 else (RED, RED_F)
    for i, w in enumerate([7, 32, 14, 16, 16, 18, 14, 12, 11, 14, 11], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.cell(ws.max_row + 2, 1, "Score = 70% Breadth% (sector ke scanned stocks me se kitne % me EMA signal) + "
            "30% Vol-confirm% (signal stocks me se kitno ka volume >= 1.5x avg)").font = Font(italic=True, color="595959")

# ===================== FETCH ALL ====================================
def fetch_all(stocks, to_date, from_date):
    out = {}
    def work(row):
        if AUTH_FAILED["flag"]:
            return row.symbol, None
        return row.symbol, fetch_daily(row.key, to_date, from_date)
    total = len(stocks); done = 0
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        for sym, df in ex.map(work, stocks.itertuples(index=False)):
            out[sym] = df
            done += 1
            if done % 250 == 0 or done == total:
                print(f"  fetched {done}/{total}")
    return out

# ===================== MAIN =========================================
def main(now=None):
    now = now or datetime.now(IST)
    today = now.date()
    warns = []
    mount_drive()

    token = get_secret(SECRET_UPSTOX)
    TG["token"], TG["chat"] = get_secret(SECRET_TG_TOKEN), get_secret(SECRET_TG_CHAT)
    if not token:
        raise SystemExit(f"Colab secret '{SECRET_UPSTOX}' nahi mila")
    HEADERS.update({"Authorization": f"Bearer {token}", "Accept": "application/json"})
    if not check_token():
        tg_send("❌ <b>EMA Scanner</b>: Upstox token invalid/expired. Naya token Colab secret me daal ke dobara run kar.")
        raise SystemExit("Upstox token invalid/expired")

    stocks = load_stocklist()
    ins = load_instruments()
    uni = stocks.merge(ins[["k", "instrument_key"]], on="k", how="left")
    unmatched = int(uni["instrument_key"].isna().sum())
    uni = uni.dropna(subset=["instrument_key"]).rename(columns={"instrument_key": "key"})
    uni = uni.drop_duplicates("symbol").drop_duplicates("key").reset_index(drop=True)   # same stock 2 baar na aaye
    print(f"Stocklist NSE symbols: {len(stocks)} | Upstox match: {len(uni)} | unmatched: {unmatched}")

    refresh = FORCE_UNIVERSE_REFRESH or not os.path.exists(UNIVERSE_CSV) or \
        (time.time() - os.path.getmtime(UNIVERSE_CSV)) > UNIVERSE_REFRESH_DAYS * 86400
    if refresh:
        scan = uni
        print(f"FULL universe scan ({len(scan)} stocks) - liquid list refresh hogi...")
    else:
        liq_syms = set(pd.read_csv(UNIVERSE_CSV)["symbol"])
        scan = uni[uni["symbol"].isin(liq_syms)].reset_index(drop=True)
        print(f"Liquid universe scan ({len(scan)} stocks)")

    from_date = (today - timedelta(days=HISTORY_DAYS)).isoformat()
    data = fetch_all(scan, today.isoformat(), from_date)
    if AUTH_FAILED["flag"]:
        tg_send("❌ <b>EMA Scanner</b>: Upstox 401 aaya (token expire). Naya token daal ke dobara run kar.")
        raise SystemExit("401 Unauthorized")
    failed = [s for s, d in data.items() if d is None or len(d) == 0]
    data = {s: d for s, d in data.items() if d is not None and len(d) > 0}
    if not data:
        tg_send("❌ <b>EMA Scanner</b>: Koi data nahi aaya. API/token check kar.")
        raise SystemExit("No data")

    def mode_date(dct):
        return pd.Series([d["date"].iloc[-1] for d in dct.values()]).mode().iloc[0]
    data_date = mode_date(data)

    # Aaj ka candle historical API me na aaye to intraday(days/1) se try
    market_closed = now.time() >= dtime(15, 45)
    if today.weekday() < 5 and market_closed and data_date < today:
        probe_sym = next(iter(data))
        probe_key = scan.loc[scan["symbol"] == probe_sym, "key"].iloc[0]
        pdf = fetch_today_intraday(probe_key)
        if pdf is not None and pdf["date"].iloc[-1] == today:
            print("Aaj ka candle historical me nahi tha -> intraday endpoint se utha raha hu...")
            keys = scan.set_index("symbol")["key"].to_dict()
            def add_today(sym):
                t = fetch_today_intraday(keys[sym])
                return sym, t
            with ThreadPoolExecutor(max_workers=WORKERS) as ex:
                for sym, t in ex.map(add_today, list(data.keys())):
                    if t is not None and t["date"].iloc[-1] > data[sym]["date"].iloc[-1]:
                        data[sym] = pd.concat([data[sym], t.tail(1)], ignore_index=True)
            data_date = mode_date(data)
    if data_date != today:
        warns.append(f"Latest data date {data_date:%d-%b-%Y} hai (aaj {today:%d-%b} ka candle nahi aaya / market band tha)")
    if today.weekday() < 5 and not market_closed and data_date == today:
        warns.append("Market band hone se pehle run hua - aaj ka candle adhoora ho sakta hai (4 PM ke baad chala)")

    # ---- analysis + liquidity filter ----
    meta = uni.set_index("symbol")
    rows, liquid_syms = [], []
    for sym, df in data.items():
        if df["date"].iloc[-1] != data_date:
            continue
        a = analyse(df)
        if a is None:
            continue
        if a["close"] >= MIN_PRICE and a["tv_cr"] >= MIN_TRADED_VALUE_CR:
            liquid_syms.append(sym)
            a.update({"symbol": sym, "name": meta.at[sym, "name"], "sector": meta.at[sym, "sector"],
                      "industry": meta.at[sym, "industry"]})
            rows.append(a)
    if refresh:
        pd.DataFrame({"symbol": liquid_syms}).to_csv(UNIVERSE_CSV, index=False)
        print(f"Liquid universe saved: {len(liquid_syms)} stocks")
    res = pd.DataFrame(rows)
    if res.empty:
        tg_send("⚠️ <b>EMA Scanner</b>: Liquidity filter ke baad koi stock nahi bacha. Config check kar.")
        raise SystemExit("empty result")
    res["spike"] = (res["vol_ratio"] >= VOL_SPIKE_X) & (res["chg_pct"] > 0)

    # ---- Signals (top N) ----
    sig = res[res["signal"]].copy()
    sort_col = "vol_ratio" if RANK_BY == "vol_ratio" else "volume"
    sig = sig.sort_values(sort_col, ascending=False).head(TOP_N).reset_index(drop=True)
    n_signals_all, n_fresh_all = int(res["signal"].sum()), int((res["signal"] & (res["cross_k"] == 0)).sum())
    sig_out = pd.DataFrame({
        "Date": data_date, "Rank": range(1, len(sig) + 1), "Symbol": sig["symbol"], "Company": sig["name"],
        "Sector": sig["sector"], "Industry": sig["industry"], "Close": sig["close"].round(2),
        "Chg %": sig["chg_pct"].round(2), "EMA5": sig["ema_f"].round(2), "EMA8": sig["ema_m"].round(2),
        "EMA13": sig["ema_s"].round(2), "Cross": sig["cross"], "Signal": sig["cross_k"].map(cross_tag),
        "EMAs below close (of 3)": sig["above_n"], "5>8>13 aligned": sig["aligned"].map({True: "Yes", False: "No"}),
        "Volume": sig["volume"].round(0), "Avg Vol 20d": sig["avg_vol"].round(0),
        "Vol Ratio (x)": sig["vol_ratio"].round(2), "Traded Value (Cr)": sig["tv_cr"].round(2)})

    # ---- Volume spike ----
    sp = res[res["spike"]].sort_values("vol_ratio", ascending=False).head(VOL_SPIKE_TOP).reset_index(drop=True)
    sig_set = set(res.loc[res["signal"], "symbol"])
    sp_out = pd.DataFrame({
        "Date": data_date, "Rank": range(1, len(sp) + 1), "Symbol": sp["symbol"], "Company": sp["name"],
        "Sector": sp["sector"], "Close": sp["close"].round(2), "Chg %": sp["chg_pct"].round(2),
        "Volume": sp["volume"].round(0), "Avg Vol 20d": sp["avg_vol"].round(0),
        "Vol Ratio (x)": sp["vol_ratio"].round(2), "Traded Value (Cr)": sp["tv_cr"].round(2),
        "In EMA Signal": sp["symbol"].map(lambda s: "Yes" if s in sig_set else "No")})

    # ---- Sector log -> heatmaps / momentum ----
    log = update_sector_log(res, data_date)

    fm = {"Close": "0.00", "Chg %": "0.00", "EMA5": "0.00", "EMA8": "0.00", "EMA13": "0.00",
          "Volume": "#,##0", "Avg Vol 20d": "#,##0", "Vol Ratio (x)": "0.00", "Traded Value (Cr)": "0.00"}
    wd = {"Company": 34, "Sector": 28, "Industry": 28, "Cross": 14, "EMAs below close (of 3)": 14}
    wb = get_wb(MASTER_XLSX)
    append_sheet(wb, "Signals", sig_out, data_date, fm, wd)
    append_sheet(wb, "Volume Spike", sp_out, data_date, fm, wd)
    write_heatmap(wb, "Sector Daily", daily_matrix(log, "Signals"), "0",
                  "Value = us din us sector me EMA signal wale stocks ki sankhya")
    write_heatmap(wb, "Sector Weekly", weekly_matrix(log, "Signals"), "0.0",
                  "Value = hafte ke daily signal count ka average")
    save_wb(wb, MASTER_XLSX, ["Signals", "Volume Spike", "Sector Daily", "Sector Weekly"])

    wb2 = get_wb(MOMENTUM_XLSX)
    write_latest_ranking(wb2, log, data_date)
    write_heatmap(wb2, "Score Daily", daily_matrix(log, "Score"), "0.0",
                  "Momentum score 0-100 (daily)")
    write_heatmap(wb2, "Score Weekly", weekly_matrix(log, "Score"), "0.0",
                  "Momentum score 0-100 (hafte ka average)")
    save_wb(wb2, MOMENTUM_XLSX, ["Latest Ranking", "Score Daily", "Score Weekly"])

    # ---- Telegram ----
    e = html.escape
    L = [f"📈 <b>EMA 5-8-13 Scanner</b> | {data_date:%d-%b-%Y}"]
    if warns:
        L += [f"⚠️ {e(w)}" for w in warns]
    L.append(f"Scanned: {len(res)} liquid | Signals: {n_signals_all} (Fresh: {n_fresh_all}) | Vol spikes: {int(res['spike'].sum())}")
    L.append("")
    rank_lbl = "Vol Ratio" if RANK_BY == "vol_ratio" else "Volume"
    L.append(f"🔥 <b>Top 10 (by {rank_lbl})</b>")
    for r in sig.head(10).itertuples():
        L.append(f"{r.Index + 1}. <b>{e(r.symbol)}</b> ₹{r.close:,.2f} ({r.chg_pct:+.1f}%) | {r.vol_ratio:.1f}x | "
                 f"{e(r.cross)} {cross_tag(r.cross_k)} | {e(r.sector)}")
    if sig.empty:
        L.append("Aaj koi signal nahi.")
    ld = log[log["Date"] == data_date].dropna(subset=["Score"]).sort_values("Score", ascending=False)
    dates = sorted(log["Date"].unique())
    prevd = dates[-2] if len(dates) > 1 else None
    prevs = log[log["Date"] == prevd].set_index("Sector")["Score"] if prevd else pd.Series(dtype=float)
    if not ld.empty:
        L += ["", "🏭 <b>Top sectors (momentum score)</b>"]
        for i, r in enumerate(ld.head(5).itertuples(), 1):
            ch = ""
            if r.Sector in prevs.index and not pd.isna(prevs[r.Sector]):
                d = r.Score - prevs[r.Sector]
                ch = f" ({'▲' if d > 0 else '▼' if d < 0 else '='}{abs(d):.1f})"
            L.append(f"{i}. {e(r.Sector)} - {r.Score:.1f}{ch} [{r.Signals}/{r.Total}]")
    if len(sp):
        L += ["", "⚡ <b>Volume spikes (top 5)</b>"]
        for r in sp.head(5).itertuples():
            L.append(f"• <b>{e(r.symbol)}</b> ₹{r.close:,.2f} ({r.chg_pct:+.1f}%) | {r.vol_ratio:.1f}x | {e(r.sector)}")
    if failed:
        L += ["", f"ℹ️ {len(failed)} stocks ka data nahi aaya"]
    tg_send("\n".join(L))
    if SEND_EXCEL_TO_TG:
        tg_doc(MASTER_XLSX, f"Master file {data_date:%d-%b-%Y}")
        tg_doc(MOMENTUM_XLSX, f"Sector momentum {data_date:%d-%b-%Y}")
    print("Done. Excel files Drive me save ho gayi:", MASTER_XLSX, "|", MOMENTUM_XLSX)
    return sig_out, sp_out, log


if __name__ == "__main__":
    main()
